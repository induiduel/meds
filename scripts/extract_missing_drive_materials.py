# -*- coding: utf-8 -*-
"""
extract_missing_drive_materials.py
Extracts text from all pending PDF, PPTX, and DOCX files in meds_database/meds_sorular
and saves them to meds_database/meds_sorular_txt.
"""

import os
import sys
import zipfile
import xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding='utf-8')

# Try importing fitz (PyMuPDF)
try:
    import fitz
    has_fitz = True
except ImportError:
    has_fitz = False

# Try importing pptx
try:
    import pptx
    has_pptx = True
except ImportError:
    has_pptx = False

SRC_DIR = ((__import__('os').environ.get('MEDS_DATABASE_DIR') or __import__('os').path.expanduser('~/meds_database')) + "/meds_sorular")
DEST_DIR = ((__import__('os').environ.get('MEDS_DATABASE_DIR') or __import__('os').path.expanduser('~/meds_database')) + "/meds_sorular_txt")
os.makedirs(DEST_DIR, exist_ok=True)

def extract_pdf(pdf_path):
    doc = fitz.open(pdf_path)
    pages = []
    for i, page in enumerate(doc):
        text = page.get_text("text").strip()
        pages.append(f"--- [SAYFA {i+1}] ---\n{text}")
    doc.close()
    return "\n\n".join(pages), len(pages)

def extract_pptx(pptx_path):
    prs = pptx.Presentation(pptx_path)
    slides = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        texts.append(line)
        slide_text = "\n".join(texts)
        slides.append(f"--- [SLAYT {i+1}] ---\n{slide_text}")
    return "\n\n".join(slides), len(slides)

def extract_docx(docx_path):
    with zipfile.ZipFile(docx_path) as z:
        xml_content = z.read("word/document.xml")
    root = ET.fromstring(xml_content)
    # Extract all text elements: w:t
    namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paragraphs = []
    for p in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
        texts = [t.text for t in p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if t.text]
        line = "".join(texts).strip()
        if line:
            paragraphs.append(line)
    return "\n\n".join(paragraphs), len(paragraphs)

print("Starting scan of meds_sorular for missing text extractions...")
processed_count = 0
error_count = 0

for root, dirs, files in os.walk(SRC_DIR):
    for file in files:
        base_name, ext = os.path.splitext(file)
        ext = ext.lower()
        if ext not in [".pdf", ".pptx", ".docx"]:
            continue
        
        src_path = os.path.join(root, file)
        dest_txt_path = os.path.join(DEST_DIR, f"{base_name}.txt")

        # Skip if already exists and size > 100 bytes
        if os.path.exists(dest_txt_path) and os.path.getsize(dest_txt_path) > 100:
            continue

        print(f"Extracting: {file}...")
        try:
            if ext == ".pdf":
                text, count = extract_pdf(src_path)
                unit = "sayfa"
            elif ext == ".pptx":
                text, count = extract_pptx(src_path)
                unit = "slayt"
            elif ext == ".docx":
                text, count = extract_docx(src_path)
                unit = "paragraf"
            
            with open(dest_txt_path, "w", encoding="utf-8") as out:
                out.write(text)
            
            print(f"  ✓ Saved to {base_name}.txt ({count} {unit}, {len(text)} chars)")
            processed_count += 1
        except Exception as e:
            print(f"  ❌ Error processing {file}: {e}")
            error_count += 1

print(f"\nProcessing complete! Processed: {processed_count}, Errors: {error_count}")
